from __future__ import annotations

import time
import warnings
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Tuple

import numpy as np
import pandas as pd
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy.stats import jarque_bera, kurtosis, skew
from statsmodels.graphics.tsaplots import plot_acf, plot_pacf
from statsmodels.stats.diagnostic import acorr_ljungbox, het_arch
from statsmodels.tsa.arima.model import ARIMA

warnings.filterwarnings("ignore")

try:
    import matplotlib.pyplot as plt
except ImportError as exc:
    raise ImportError("matplotlib não está instalado.") from exc

try:
    from arch import arch_model
except ImportError as exc:
    raise ImportError(
        "A biblioteca 'arch' não está instalada. Instale com: pip install arch"
    ) from exc


# =====================================================================
# 1) CONFIGURAÇÕES
# =====================================================================

@dataclass
class Configuracao:
    """Parâmetros principais da coleta e da análise."""

    # Coleta
    data_inicial: str = "2002-01-01"
    data_final: str = "2026-04-01"
    tamanho_bloco_dias: int = 180
    timeout: int = 30
    max_tentativas: int = 3
    pausa_entre_requisicoes: float = 0.4

    # Arquivos de saída
    salvar_base_excel: bool = True
    arquivo_base_excel: str = "dados_usdbrl_bcb.xlsx"
    aba_base_excel: str = "usd_brl"

    pasta_saida: str = "resultados_garch_usdbrl"
    pasta_figuras: str = "figuras"
    pasta_tabelas: str = "tabelas"

    # Colunas da base
    coluna_data: str = "data"
    coluna_preco: str = "cotacaoVenda"
    coluna_retorno: str = "retorno_log"

    # Modelagem
    lags_correlograma: int = 30
    lags_ljung_box: Tuple[int, ...] = (10, 20, 30)
    lags_arch_lm: int = 10
    criterio_escolha_media: str = "bic"
    nivel_significancia: float = 0.05
    candidatos_media: Tuple[Tuple[int, int], ...] = (
        (0, 0),
        (1, 0),
        (0, 1),
        (1, 1),
        (2, 0),
        (0, 2),
    )


CFG = Configuracao()


# =====================================================================
# 2) PADRÃO VISUAL
# =====================================================================

COR_PRINCIPAL = "#1f4e79"
COR_SECUNDARIA = "#7f8c8d"
COR_GRADE = "#d9d9d9"
COR_CABECALHO = "#dbe5f1"
COR_BORDA = "#bfbfbf"


def configurar_estilo_graficos() -> None:
    """Define o estilo básico dos gráficos."""
    plt.rcParams.update(
        {
            "figure.dpi": 150,
            "savefig.dpi": 120,
            "figure.figsize": (12, 5),
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "axes.titleweight": "bold",
            "axes.edgecolor": "#666666",
            "axes.linewidth": 0.8,
            "font.size": 11,
            "grid.color": COR_GRADE,
            "grid.linestyle": "--",
            "grid.linewidth": 0.7,
            "grid.alpha": 0.8,
            "legend.frameon": False,
            "xtick.color": "#333333",
            "ytick.color": "#333333",
        }
    )


# =====================================================================
# 3) COLETA DOS DADOS PTAX
# =====================================================================


def montar_url_ptax(inicio: datetime, fim: datetime) -> str:
    """Monta a URL de consulta da API PTAX para um intervalo de datas."""
    inicio_fmt = inicio.strftime("%m-%d-%Y")
    fim_fmt = fim.strftime("%m-%d-%Y")

    return (
        "https://olinda.bcb.gov.br/olinda/servico/PTAX/versao/v1/odata/"
        "CotacaoDolarPeriodo(dataInicial=@dataInicial,dataFinalCotacao=@dataFinalCotacao)?"
        f"@dataInicial='{inicio_fmt}'&@dataFinalCotacao='{fim_fmt}'"
        "&$top=10000"
        "&$format=json"
        "&$select=cotacaoCompra,cotacaoVenda,dataHoraCotacao"
    )


def baixar_bloco(inicio: datetime, fim: datetime, cfg: Configuracao) -> pd.DataFrame:
    """Baixa um bloco de observações da API PTAX."""
    url = montar_url_ptax(inicio, fim)
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Accept": "application/json",
    }

    ultima_excecao = None
    for tentativa in range(1, cfg.max_tentativas + 1):
        try:
            resposta = requests.get(url, headers=headers, timeout=cfg.timeout)
            resposta.raise_for_status()
            dados = resposta.json()
            return pd.DataFrame(dados.get("value", []))
        except Exception as exc:
            ultima_excecao = exc
            if tentativa < cfg.max_tentativas:
                time.sleep(1.5 * tentativa)
            else:
                raise RuntimeError(
                    f"Falha ao baixar o bloco {inicio:%Y-%m-%d} até {fim:%Y-%m-%d}."
                ) from ultima_excecao

    raise RuntimeError("Erro inesperado na rotina de download.")


def gerar_blocos(data_inicial: datetime, data_final: datetime, tamanho_bloco_dias: int):
    """Gera os intervalos de consulta usados no download."""
    atual = data_inicial
    while atual <= data_final:
        fim_bloco = min(atual + timedelta(days=tamanho_bloco_dias - 1), data_final)
        yield atual, fim_bloco
        atual = fim_bloco + timedelta(days=1)


def tratar_dados(df: pd.DataFrame) -> pd.DataFrame:
    """Organiza a base e calcula as variáveis usadas na análise."""
    if df.empty:
        raise ValueError("Nenhum dado foi retornado pela API do Banco Central.")

    base = df.copy()
    base["dataHoraCotacao"] = pd.to_datetime(base["dataHoraCotacao"], errors="coerce")
    base["cotacaoCompra"] = pd.to_numeric(base["cotacaoCompra"], errors="coerce")
    base["cotacaoVenda"] = pd.to_numeric(base["cotacaoVenda"], errors="coerce")

    base = base.dropna(subset=["dataHoraCotacao", "cotacaoVenda"])
    base = base.sort_values("dataHoraCotacao")

    base["data"] = base["dataHoraCotacao"].dt.normalize()
    base = base.drop_duplicates(subset="data", keep="last")

    base["log_cotacao_venda"] = np.log(base["cotacaoVenda"])
    base["retorno_log"] = 100 * base["log_cotacao_venda"].diff()
    base["retorno_percentual"] = base["cotacaoVenda"].pct_change() * 100

    colunas_finais = [
        "data",
        "dataHoraCotacao",
        "cotacaoCompra",
        "cotacaoVenda",
        "log_cotacao_venda",
        "retorno_log",
        "retorno_percentual",
    ]
    return base[colunas_finais].reset_index(drop=True)


def coletar_base_ptax(cfg: Configuracao) -> pd.DataFrame:
    """Executa a coleta completa da série PTAX."""
    data_inicial = datetime.strptime(cfg.data_inicial, "%Y-%m-%d")
    data_final = datetime.strptime(cfg.data_final, "%Y-%m-%d")

    if data_inicial > data_final:
        raise ValueError("data_inicial não pode ser maior que data_final.")

    partes = []
    for inicio, fim in gerar_blocos(data_inicial, data_final, cfg.tamanho_bloco_dias):
        parte = baixar_bloco(inicio, fim, cfg)
        partes.append(parte)
        time.sleep(cfg.pausa_entre_requisicoes)

    bruto = pd.concat(partes, ignore_index=True) if partes else pd.DataFrame()
    return tratar_dados(bruto)


# =====================================================================
# 4) UTILITÁRIOS DE ARQUIVO
# =====================================================================


def criar_pastas_saida(cfg: Configuracao) -> Dict[str, Path]:
    """Cria as pastas de saída."""
    pasta_raiz = Path(cfg.pasta_saida)
    pasta_figuras = pasta_raiz / cfg.pasta_figuras
    pasta_tabelas = pasta_raiz / cfg.pasta_tabelas

    pasta_raiz.mkdir(parents=True, exist_ok=True)
    pasta_figuras.mkdir(parents=True, exist_ok=True)
    pasta_tabelas.mkdir(parents=True, exist_ok=True)

    return {
        "raiz": pasta_raiz,
        "figuras": pasta_figuras,
        "tabelas": pasta_tabelas,
    }


def salvar_em_excel(df: pd.DataFrame, caminho_arquivo: Path, nome_aba: str) -> None:
    """Salva a base em Excel e ajusta a largura das colunas."""
    with pd.ExcelWriter(caminho_arquivo, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=nome_aba, index=False)
        worksheet = writer.sheets[nome_aba]

        for col in worksheet.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                valor = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(valor))
            worksheet.column_dimensions[col_letter].width = max_length + 2


def arredondar_tabela(df: pd.DataFrame, casas: int = 6) -> pd.DataFrame:
    """Arredonda as colunas numéricas."""
    tabela = df.copy()
    colunas_numericas = tabela.select_dtypes(include=[np.number]).columns
    tabela[colunas_numericas] = tabela[colunas_numericas].round(casas)
    return tabela


def formatar_planilha_openpyxl(worksheet) -> None:
    """Aplica uma formatação simples à planilha exportada."""
    preenchimento = PatternFill(fill_type="solid", fgColor=COR_CABECALHO.replace("#", ""))
    borda = Border(
        left=Side(style="thin", color=COR_BORDA.replace("#", "")),
        right=Side(style="thin", color=COR_BORDA.replace("#", "")),
        top=Side(style="thin", color=COR_BORDA.replace("#", "")),
        bottom=Side(style="thin", color=COR_BORDA.replace("#", "")),
    )

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = preenchimento
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = borda

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = borda
            cell.alignment = Alignment(vertical="center")

    worksheet.freeze_panes = "A2"

    for coluna in worksheet.columns:
        max_length = 0
        letra = coluna[0].column_letter
        for cell in coluna:
            valor = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(valor))
        worksheet.column_dimensions[letra].width = min(max_length + 2, 40)


def nome_aba_excel_unico(nome_original: str, nomes_usados: set[str]) -> str:
    """Gera um nome de aba único dentro do limite do Excel."""
    base = nome_original[:31]
    if base not in nomes_usados:
        nomes_usados.add(base)
        return base

    contador = 2
    while True:
        sufixo = f"_{contador}"
        candidato = f"{base[:31 - len(sufixo)]}{sufixo}"
        if candidato not in nomes_usados:
            nomes_usados.add(candidato)
            return candidato
        contador += 1


def exportar_tabelas_excel(tabelas: Dict[str, pd.DataFrame], caminho_saida: Path) -> None:
    """Exporta as tabelas para um único arquivo Excel."""
    nomes_usados: set[str] = set()
    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        for nome_aba, tabela in tabelas.items():
            nome_aba_excel = nome_aba_excel_unico(nome_aba, nomes_usados)
            arredondada = arredondar_tabela(tabela)
            arredondada.to_excel(writer, sheet_name=nome_aba_excel, index=False)
            formatar_planilha_openpyxl(writer.sheets[nome_aba_excel])


def exportar_tabelas_csv(tabelas: Dict[str, pd.DataFrame], pasta_tabelas: Path) -> None:
    """Salva cada tabela em CSV."""
    for nome, tabela in tabelas.items():
        arredondar_tabela(tabela).to_csv(
            pasta_tabelas / f"{nome}.csv",
            index=False,
            encoding="utf-8-sig",
        )


# =====================================================================
# 5) PREPARAÇÃO DA BASE PARA MODELAGEM
# =====================================================================


def preparar_base_modelagem(df: pd.DataFrame, cfg: Configuracao) -> pd.DataFrame:
    """Prepara a base que será usada na modelagem."""
    base = df.copy()
    base[cfg.coluna_data] = pd.to_datetime(base[cfg.coluna_data], errors="coerce")
    base[cfg.coluna_preco] = pd.to_numeric(base[cfg.coluna_preco], errors="coerce")
    base[cfg.coluna_retorno] = pd.to_numeric(base[cfg.coluna_retorno], errors="coerce")

    base = base.dropna(subset=[cfg.coluna_data, cfg.coluna_preco])
    base = base.sort_values(cfg.coluna_data).reset_index(drop=True)
    base["retorno_quadrado"] = base[cfg.coluna_retorno] ** 2

    return base.dropna(subset=[cfg.coluna_retorno]).copy().reset_index(drop=True)


# =====================================================================
# 6) ESTATÍSTICAS DESCRITIVAS E TESTES PRELIMINARES
# =====================================================================


def estatisticas_descritivas(serie: pd.Series, nome: str) -> pd.DataFrame:
    """Calcula estatísticas descritivas da série."""
    s = pd.Series(serie).dropna().astype(float)
    jb = jarque_bera(s)

    return pd.DataFrame(
        {
            "serie": [nome],
            "n": [int(s.size)],
            "media": [s.mean()],
            "mediana": [s.median()],
            "desvio_padrao": [s.std(ddof=1)],
            "minimo": [s.min()],
            "q1": [s.quantile(0.25)],
            "q3": [s.quantile(0.75)],
            "maximo": [s.max()],
            "assimetria": [skew(s, bias=False)],
            "curtose_excesso": [kurtosis(s, fisher=True, bias=False)],
            "jarque_bera": [jb.statistic],
            "jb_pvalor": [jb.pvalue],
        }
    )


def teste_ljung_box(serie: pd.Series, lags: Tuple[int, ...]) -> pd.DataFrame:
    """Aplica o teste de Ljung-Box."""
    resultado = acorr_ljungbox(serie.dropna(), lags=list(lags), return_df=True)
    return resultado.reset_index().rename(
        columns={"index": "lag", "lb_stat": "estatistica_lb", "lb_pvalue": "pvalor_lb"}
    )


def teste_arch_lm(serie: pd.Series, lags: int) -> pd.DataFrame:
    """Aplica o teste ARCH-LM."""
    lm_stat, lm_pvalor, f_stat, f_pvalor = het_arch(serie.dropna(), nlags=lags)
    return pd.DataFrame(
        {
            "lags": [lags],
            "lm_stat": [lm_stat],
            "lm_pvalor": [lm_pvalor],
            "f_stat": [f_stat],
            "f_pvalor": [f_pvalor],
        }
    )


# =====================================================================
# 7) GRÁFICOS
# =====================================================================


def salvar_figura(fig, caminho: Path) -> None:
    """Salva a figura no caminho indicado."""
    fig.tight_layout()
    fig.savefig(caminho.with_suffix(".pdf"), bbox_inches="tight")
    plt.close(fig)


def formatar_eixo_data(ax, datas: pd.Series) -> None:
    """Mostra apenas os anos no eixo x e preserva a última data completa."""
    serie_datas = pd.to_datetime(pd.Series(datas).dropna())
    if serie_datas.empty:
        return

    data_inicial = serie_datas.min().normalize()
    data_final = serie_datas.max().normalize()

    anos = list(range(data_inicial.year, data_final.year, 2))
    ticks = [pd.Timestamp(year=ano, month=1, day=1) for ano in anos]
    ticks = [tick for tick in ticks if data_inicial <= tick < data_final]
    ticks.append(data_final)

    labels = [str(tick.year) for tick in ticks[:-1]] + [data_final.strftime("%d/%m/%Y")]

    ax.set_xlim(data_inicial, data_final)
    ax.set_xticks(ticks)
    ax.set_xticklabels(labels, rotation=45, ha="right")


def grafico_serie_nivel(base: pd.DataFrame, cfg: Configuracao, pasta_figuras: Path) -> None:
    """Gera o gráfico da série em nível."""
    fig, ax = plt.subplots(figsize=(13, 5))
    ax.plot(base[cfg.coluna_data], base[cfg.coluna_preco], color=COR_PRINCIPAL, linewidth=1.2)
    ax.set_title("USD/BRL em nível - cotação de venda")
    ax.set_xlabel("Data")
    ax.set_ylabel("Cotação")
    ax.grid(True)
    formatar_eixo_data(ax, base[cfg.coluna_data])
    salvar_figura(fig, pasta_figuras / "01_serie_em_nivel.png")


def grafico_retorno(base: pd.DataFrame, cfg: Configuracao, pasta_figuras: Path) -> None:
    """Gera o gráfico da série de retornos."""
    fig, ax = plt.subplots(figsize=(13, 5))
    ax.plot(base[cfg.coluna_data], base[cfg.coluna_retorno], color=COR_PRINCIPAL, linewidth=0.9)
    ax.axhline(0, color=COR_SECUNDARIA, linewidth=1.0)
    ax.set_title("Retornos logarítmicos diários (%)")
    ax.set_xlabel("Data")
    ax.set_ylabel("Retorno (%)")
    ax.grid(True)
    formatar_eixo_data(ax, base[cfg.coluna_data])
    salvar_figura(fig, pasta_figuras / "02_retorno_logaritmico_pct.png")


def grafico_acf_pacf_retorno(base: pd.DataFrame, cfg: Configuracao, pasta_figuras: Path) -> None:
    """Gera os correlogramas dos retornos."""
    fig, axes = plt.subplots(2, 1, figsize=(11, 8))
    plot_acf(base[cfg.coluna_retorno], lags=cfg.lags_correlograma, ax=axes[0], zero=False)
    axes[0].set_title("FAC dos retornos logarítmicos (%)")
    axes[0].grid(True)

    plot_pacf(
        base[cfg.coluna_retorno],
        lags=cfg.lags_correlograma,
        ax=axes[1],
        zero=False,
        method="ywm",
    )
    axes[1].set_title("FACP dos retornos logarítmicos (%)")
    axes[1].grid(True)

    salvar_figura(fig, pasta_figuras / "03_fac_fapc_retorno.png")


def grafico_acf_retorno_quadrado(base: pd.DataFrame, cfg: Configuracao, pasta_figuras: Path) -> None:
    """Gera a FAC dos retornos ao quadrado."""
    fig, ax = plt.subplots(figsize=(11, 5))
    plot_acf(base["retorno_quadrado"], lags=cfg.lags_correlograma, ax=ax, zero=False)
    ax.set_title("FAC dos retornos ao quadrado")
    ax.grid(True)
    salvar_figura(fig, pasta_figuras / "04_fac_retorno_quadrado.png")


def graficos_diagnostico(
    datas: pd.Series,
    residuos_padronizados: pd.Series,
    volatilidade_condicional: pd.Series,
    cfg: Configuracao,
    pasta_figuras: Path,
) -> None:
    """Gera os gráficos de diagnóstico do ajuste."""
    fig, axes = plt.subplots(2, 1, figsize=(13, 8), sharex=True)

    axes[0].plot(datas, residuos_padronizados, color=COR_PRINCIPAL, linewidth=0.8)
    axes[0].axhline(0, color=COR_SECUNDARIA, linewidth=1.0)
    axes[0].set_title("Resíduos padronizados")
    axes[0].set_ylabel("ẑt")
    axes[0].grid(True)

    axes[1].plot(datas, volatilidade_condicional, color=COR_PRINCIPAL, linewidth=0.9)
    axes[1].set_title("Volatilidade condicional estimada")
    axes[1].set_xlabel("Data")
    axes[1].set_ylabel("σt")
    axes[1].grid(True)

    formatar_eixo_data(axes[1], datas)
    salvar_figura(fig, pasta_figuras / "05_residuos_e_volatilidade.png")

    fig, axes = plt.subplots(2, 1, figsize=(11, 8))
    plot_acf(residuos_padronizados, lags=cfg.lags_correlograma, ax=axes[0], zero=False)
    axes[0].set_title("FAC dos resíduos padronizados")
    axes[0].grid(True)

    plot_acf(residuos_padronizados**2, lags=cfg.lags_correlograma, ax=axes[1], zero=False)
    axes[1].set_title("FAC dos resíduos padronizados ao quadrado")
    axes[1].grid(True)

    salvar_figura(fig, pasta_figuras / "06_fac_residuos_padronizados.png")


# =====================================================================
# 8) ESCOLHA DA EQUAÇÃO DA MÉDIA
# =====================================================================


def nome_modelo_media(p: int, q: int) -> str:
    """Monta o nome do modelo de média."""
    if p == 0 and q == 0:
        return "Constante"
    if p > 0 and q == 0:
        return f"AR({p})"
    if p == 0 and q > 0:
        return f"MA({q})"
    return f"ARMA({p},{q})"


def ajustar_candidatos_media(
    serie_retorno: pd.Series,
    cfg: Configuracao,
) -> Tuple[pd.DataFrame, Dict[str, object], str]:
    """Ajusta os modelos candidatos para a equação da média."""
    s = serie_retorno.dropna().astype(float)
    resultados = []
    objetos_ajustados: Dict[str, object] = {}

    lb_bruto = teste_ljung_box(s, cfg.lags_ljung_box)
    retornos_sem_dependencia_linear_forte = bool(
        (lb_bruto["pvalor_lb"] > cfg.nivel_significancia).all()
    )

    for p, q in cfg.candidatos_media:
        nome = nome_modelo_media(p, q)
        try:
            modelo = ARIMA(
                s,
                order=(p, 0, q),
                trend="c",
                enforce_stationarity=False,
                enforce_invertibility=False,
            )
            ajuste = modelo.fit()
            residuos = pd.Series(ajuste.resid).dropna()
            lb_resid = teste_ljung_box(residuos, cfg.lags_ljung_box)

            resultados.append(
                {
                    "modelo_media": nome,
                    "p": p,
                    "q": q,
                    "aic": ajuste.aic,
                    "bic": ajuste.bic,
                    "loglik": ajuste.llf,
                    "n_parametros": int(len(ajuste.params)),
                    "lb_pvalor_lag10": float(lb_resid.loc[lb_resid["lag"] == 10, "pvalor_lb"].iloc[0]),
                    "lb_pvalor_lag20": float(lb_resid.loc[lb_resid["lag"] == 20, "pvalor_lb"].iloc[0]),
                    "lb_pvalor_lag30": float(lb_resid.loc[lb_resid["lag"] == 30, "pvalor_lb"].iloc[0]),
                    "convergiu": True,
                }
            )
            objetos_ajustados[nome] = ajuste
        except Exception as exc:
            resultados.append(
                {
                    "modelo_media": nome,
                    "p": p,
                    "q": q,
                    "aic": np.nan,
                    "bic": np.nan,
                    "loglik": np.nan,
                    "n_parametros": np.nan,
                    "lb_pvalor_lag10": np.nan,
                    "lb_pvalor_lag20": np.nan,
                    "lb_pvalor_lag30": np.nan,
                    "convergiu": False,
                    "erro": str(exc),
                }
            )

    tabela = pd.DataFrame(resultados)
    criterio = cfg.criterio_escolha_media.lower()
    if criterio not in {"aic", "bic"}:
        raise ValueError("criterio_escolha_media deve ser 'aic' ou 'bic'.")

    if retornos_sem_dependencia_linear_forte and "Constante" in objetos_ajustados:
        modelo_escolhido = "Constante"
    else:
        modelos_validos = tabela[
            (tabela["convergiu"] == True)
            & (tabela["lb_pvalor_lag10"] > cfg.nivel_significancia)
            & (tabela["lb_pvalor_lag20"] > cfg.nivel_significancia)
            & (tabela["lb_pvalor_lag30"] > cfg.nivel_significancia)
        ].copy()

        if modelos_validos.empty:
            modelos_validos = tabela[tabela["convergiu"] == True].copy()

        if modelos_validos.empty:
            raise RuntimeError("Nenhum modelo de média convergiu.")

        modelo_escolhido = (
            modelos_validos.sort_values(by=[criterio, "n_parametros"], ascending=[True, True])
            .iloc[0]["modelo_media"]
        )

    return (
        tabela.sort_values(by=["bic", "aic"], na_position="last"),
        objetos_ajustados,
        modelo_escolhido,
    )


def rotular_parametro_media(nome_parametro: object) -> str:
    """Padroniza os nomes dos parâmetros da média."""
    nome = str(nome_parametro)
    mapa = {
        "const": "Constante",
        "ma.L1": "MA(1)",
        "ma.L2": "MA(2)",
        "ar.L1": "AR(1)",
        "ar.L2": "AR(2)",
        "sigma2": "Variância do erro da média (sigma2)",
    }
    return mapa.get(nome, nome)


def tabela_parametros_media(ajuste_media, nome_modelo: str) -> pd.DataFrame:
    """Organiza os parâmetros do modelo de média escolhido."""
    params = ajuste_media.params
    bse = ajuste_media.bse
    pvalues = ajuste_media.pvalues
    tvalues = ajuste_media.tvalues
    parametros = list(params.index) if hasattr(params, "index") else list(range(len(params)))

    return pd.DataFrame(
        {
            "modelo_media": nome_modelo,
            "parametro": parametros,
            "parametro_relatorio": [rotular_parametro_media(p) for p in parametros],
            "estimativa": np.asarray(params),
            "erro_padrao": np.asarray(bse),
            "estatistica_t": np.asarray(tvalues),
            "pvalor": np.asarray(pvalues),
        }
    )


# =====================================================================
# 9) ESTIMAÇÃO GARCH(1,1)
# =====================================================================


def estimar_garch_11(residuos_media: pd.Series):
    """Ajusta um GARCH(1,1) gaussiano aos resíduos da média."""
    residuos = pd.Series(residuos_media).dropna().astype(float)
    modelo = arch_model(
        residuos,
        mean="Zero",
        vol="GARCH",
        p=1,
        q=1,
        dist="normal",
        rescale=False,
    )
    return modelo.fit(disp="off", show_warning=False, cov_type="robust")


def tabela_parametros_garch(ajuste_garch) -> pd.DataFrame:
    """Organiza os parâmetros estimados do GARCH."""
    return pd.DataFrame(
        {
            "parametro": ajuste_garch.params.index,
            "estimativa": ajuste_garch.params.values,
            "erro_padrao": ajuste_garch.std_err.values,
            "estatistica_t": ajuste_garch.tvalues.values,
            "pvalor": ajuste_garch.pvalues.values,
        }
    )


def calcular_meia_vida_garch(alpha: float, beta: float) -> float:
    """Calcula a meia-vida aproximada dos choques de volatilidade."""
    persistencia = alpha + beta
    if not np.isfinite(persistencia):
        return np.nan
    if persistencia <= 0 or persistencia >= 1:
        return np.nan
    return float(np.log(0.5) / np.log(persistencia))


def medidas_resumo_garch(ajuste_garch) -> pd.DataFrame:
    """Monta um resumo do ajuste GARCH(1,1)."""
    omega = float(ajuste_garch.params.get("omega", np.nan))
    alpha = float(ajuste_garch.params.get("alpha[1]", np.nan))
    beta = float(ajuste_garch.params.get("beta[1]", np.nan))
    persistencia = alpha + beta
    variancia_incondicional = np.nan
    meia_vida = calcular_meia_vida_garch(alpha, beta)

    if np.isfinite(persistencia) and persistencia < 1:
        variancia_incondicional = omega / (1 - persistencia)

    return pd.DataFrame(
        {
            "loglik": [ajuste_garch.loglikelihood],
            "aic": [ajuste_garch.aic],
            "bic": [ajuste_garch.bic],
            "omega": [omega],
            "alpha1": [alpha],
            "beta1": [beta],
            "alpha_mais_beta": [persistencia],
            "variancia_incondicional": [variancia_incondicional],
            "meia_vida_aproximada": [meia_vida],
        }
    )


# =====================================================================
# 10) DIAGNÓSTICO DO AJUSTE
# =====================================================================


def tabela_diagnostico_consolidada(
    lb_resid: pd.DataFrame,
    lb_resid2: pd.DataFrame,
    arch_lm: pd.DataFrame,
) -> pd.DataFrame:
    """Reúne os principais testes de diagnóstico em uma tabela."""
    linhas = []

    for _, row in lb_resid.iterrows():
        linhas.append(
            {
                "serie_testada": "Resíduos padronizados",
                "teste": "Ljung-Box",
                "defasagem": int(row["lag"]),
                "estatistica": float(row["estatistica_lb"]),
                "pvalor": float(row["pvalor_lb"]),
            }
        )

    for _, row in lb_resid2.iterrows():
        linhas.append(
            {
                "serie_testada": "Quadrados dos resíduos padronizados",
                "teste": "Ljung-Box",
                "defasagem": int(row["lag"]),
                "estatistica": float(row["estatistica_lb"]),
                "pvalor": float(row["pvalor_lb"]),
            }
        )

    linha_arch = arch_lm.iloc[0]
    linhas.append(
        {
            "serie_testada": "Resíduos padronizados",
            "teste": "ARCH-LM",
            "defasagem": int(linha_arch["lags"]),
            "estatistica": float(linha_arch["lm_stat"]),
            "pvalor": float(linha_arch["lm_pvalor"]),
        }
    )

    return pd.DataFrame(linhas)


def diagnosticos_modelo(ajuste_garch, cfg: Configuracao) -> Dict[str, pd.DataFrame | pd.Series]:
    """Calcula resíduos padronizados e tabelas de diagnóstico."""
    residuos_padronizados = pd.Series(ajuste_garch.std_resid).dropna()
    volatilidade_condicional = pd.Series(ajuste_garch.conditional_volatility).dropna()

    lb_resid = teste_ljung_box(residuos_padronizados, cfg.lags_ljung_box)
    lb_resid["serie_testada"] = "residuos_padronizados"

    lb_resid2 = teste_ljung_box(residuos_padronizados**2, cfg.lags_ljung_box)
    lb_resid2["serie_testada"] = "residuos_padronizados_quadrado"

    arch_lm = teste_arch_lm(residuos_padronizados, cfg.lags_arch_lm)
    arch_lm["serie_testada"] = "residuos_padronizados"

    diagnostico_consolidado = tabela_diagnostico_consolidada(lb_resid, lb_resid2, arch_lm)

    return {
        "residuos_padronizados": residuos_padronizados,
        "volatilidade_condicional": volatilidade_condicional,
        "ljung_box_residuos": lb_resid,
        "ljung_box_residuos_quadrado": lb_resid2,
        "arch_lm_residuos": arch_lm,
        "diagnostico_consolidado": diagnostico_consolidado,
    }


# =====================================================================
# 11) RESUMO TEXTUAL
# =====================================================================


def escrever_resumo_textual(
    caminho_saida: Path,
    base: pd.DataFrame,
    modelo_media_escolhido: str,
    tabela_garch_resumo: pd.DataFrame,
    lb_preliminar_retorno: pd.DataFrame,
    lb_preliminar_retorno2: pd.DataFrame,
    diagnosticos: Dict[str, pd.DataFrame | pd.Series],
) -> None:
    """Escreve um resumo textual da execução."""
    resumo_garch = tabela_garch_resumo.iloc[0]
    lb_resid = diagnosticos["ljung_box_residuos"]
    lb_resid2 = diagnosticos["ljung_box_residuos_quadrado"]
    arch_lm = diagnosticos["arch_lm_residuos"].iloc[0]

    meia_vida_texto = (
        f"{resumo_garch['meia_vida_aproximada']:.2f}"
        if np.isfinite(resumo_garch["meia_vida_aproximada"])
        else "não definida"
    )

    texto = f"""
ANÁLISE GARCH USD/BRL - RESUMO DA EXECUÇÃO
=========================================

Número de observações na base de modelagem: {len(base)}
Período da amostra: {base['data'].min().date()} até {base['data'].max().date()}

SÉRIE MODELADA
--------------
Retorno usado na modelagem: retorno_log

ETAPA PRELIMINAR
----------------
Ljung-Box dos retornos (média):
{lb_preliminar_retorno.to_string(index=False)}

Ljung-Box dos retornos ao quadrado:
{lb_preliminar_retorno2.to_string(index=False)}

EQUAÇÃO DA MÉDIA ESCOLHIDA
--------------------------
Modelo selecionado: {modelo_media_escolhido}

MODELO GARCH(1,1)
-----------------
Log-verossimilhança: {resumo_garch['loglik']:.6f}
AIC: {resumo_garch['aic']:.6f}
BIC: {resumo_garch['bic']:.6f}
omega: {resumo_garch['omega']:.6f}
alpha[1]: {resumo_garch['alpha1']:.6f}
beta[1]: {resumo_garch['beta1']:.6f}
alpha + beta: {resumo_garch['alpha_mais_beta']:.6f}
Variância incondicional: {resumo_garch['variancia_incondicional']:.6f}
Meia-vida aproximada: {meia_vida_texto} dias úteis

DIAGNÓSTICO DOS RESÍDUOS PADRONIZADOS
-------------------------------------
Ljung-Box dos resíduos padronizados:
{lb_resid.to_string(index=False)}

Ljung-Box dos resíduos padronizados ao quadrado:
{lb_resid2.to_string(index=False)}

ARCH-LM dos resíduos padronizados:
{arch_lm.to_string()}
""".strip()

    caminho_saida.write_text(texto, encoding="utf-8")


# =====================================================================
# 12) PIPELINE PRINCIPAL
# =====================================================================


def executar_analise(base_coletada: pd.DataFrame, cfg: Configuracao, caminhos: Dict[str, Path]) -> None:
    """Executa a análise GARCH a partir da base já coletada."""
    base = preparar_base_modelagem(base_coletada, cfg)

    tabela_desc_preco = estatisticas_descritivas(base[cfg.coluna_preco], "cotacao_venda")
    tabela_desc_retorno = estatisticas_descritivas(base[cfg.coluna_retorno], "retorno_log")
    tabela_estatisticas = pd.concat([tabela_desc_preco, tabela_desc_retorno], ignore_index=True)

    tabela_lb_retorno = teste_ljung_box(base[cfg.coluna_retorno], cfg.lags_ljung_box)
    tabela_lb_retorno["serie_testada"] = "retorno_log"

    tabela_lb_retorno2 = teste_ljung_box(base["retorno_quadrado"], cfg.lags_ljung_box)
    tabela_lb_retorno2["serie_testada"] = "retorno_quadrado"

    tabela_arch_preliminar = teste_arch_lm(base[cfg.coluna_retorno], cfg.lags_arch_lm)
    tabela_arch_preliminar["serie_testada"] = "retorno_log"

    grafico_serie_nivel(base, cfg, caminhos["figuras"])
    grafico_retorno(base, cfg, caminhos["figuras"])
    grafico_acf_pacf_retorno(base, cfg, caminhos["figuras"])
    grafico_acf_retorno_quadrado(base, cfg, caminhos["figuras"])

    tabela_modelos_media, ajustes_media, nome_media_escolhida = ajustar_candidatos_media(
        base[cfg.coluna_retorno], cfg
    )
    ajuste_media_escolhido = ajustes_media[nome_media_escolhida]
    tabela_param_media = tabela_parametros_media(ajuste_media_escolhido, nome_media_escolhida)

    residuos_media = pd.Series(ajuste_media_escolhido.resid).dropna()
    base_garch = base.loc[residuos_media.index].copy()
    base_garch["residuo_media"] = residuos_media.values

    ajuste_garch = estimar_garch_11(base_garch["residuo_media"])
    tabela_param_garch = tabela_parametros_garch(ajuste_garch)
    tabela_resumo_garch = medidas_resumo_garch(ajuste_garch)
    diagnosticos = diagnosticos_modelo(ajuste_garch, cfg)

    residuos_padronizados = pd.Series(diagnosticos["residuos_padronizados"])
    volatilidade_condicional = pd.Series(diagnosticos["volatilidade_condicional"])
    datas_diag = base_garch.loc[residuos_padronizados.index, cfg.coluna_data]

    graficos_diagnostico(
        datas=datas_diag,
        residuos_padronizados=residuos_padronizados,
        volatilidade_condicional=volatilidade_condicional,
        cfg=cfg,
        pasta_figuras=caminhos["figuras"],
    )

    tabelas = {
        "base_modelagem": base,
        "estatisticas_descritivas": tabela_estatisticas,
        "ljung_box_preliminar_retorno": tabela_lb_retorno,
        "ljung_box_preliminar_retorno2": tabela_lb_retorno2,
        "arch_lm_preliminar": tabela_arch_preliminar,
        "comparacao_modelos_media": tabela_modelos_media,
        "parametros_media_escolhida": tabela_param_media,
        "parametros_garch": tabela_param_garch,
        "resumo_garch": tabela_resumo_garch,
        "lb_resid_padronizados": diagnosticos["ljung_box_residuos"],
        "lb_resid_padronizados_quad": diagnosticos["ljung_box_residuos_quadrado"],
        "arch_lm_resid_padronizados": diagnosticos["arch_lm_residuos"],
        "diagnostico_consolidado": diagnosticos["diagnostico_consolidado"],
    }

    exportar_tabelas_csv(tabelas, caminhos["tabelas"])
    exportar_tabelas_excel(tabelas, caminhos["raiz"] / "tabelas_resultados_garch.xlsx")
    escrever_resumo_textual(
        caminho_saida=caminhos["raiz"] / "resumo_execucao.txt",
        base=base,
        modelo_media_escolhido=nome_media_escolhida,
        tabela_garch_resumo=tabela_resumo_garch,
        lb_preliminar_retorno=tabela_lb_retorno,
        lb_preliminar_retorno2=tabela_lb_retorno2,
        diagnosticos=diagnosticos,
    )



def main() -> None:
    """Executa a coleta, o tratamento e a análise completa em um único script."""
    configurar_estilo_graficos()
    caminhos = criar_pastas_saida(CFG)

    base_coletada = coletar_base_ptax(CFG)

    if CFG.salvar_base_excel:
        caminho_excel = Path(CFG.arquivo_base_excel)
        salvar_em_excel(base_coletada, caminho_excel, CFG.aba_base_excel)

    executar_analise(base_coletada, CFG, caminhos)


if __name__ == "__main__":
    main()
